import io
import PyPDF2
from docx import Document
from openpyxl import load_workbook
from PIL import Image
import pytesseract

class FileHandler:
    async def read_file(self, file):
        content = await file.read()
        if file.filename.endswith('.pdf'):
            return self.read_pdf(io.BytesIO(content))
        elif file.filename.endswith('.docx'):
            return self.read_docx(io.BytesIO(content))
        elif file.filename.endswith('.xlsx'):
            return self.read_excel(io.BytesIO(content))
        elif file.filename.lower().endswith(('.png', '.jpg', '.jpeg')):
            return self.extract_text_from_image(Image.open(io.BytesIO(content)))
        else:
            return content.decode()

    def read_pdf(self, file):
        reader = PyPDF2.PdfReader(file)
        return "\n".join([page.extract_text() for page in reader.pages])

    def read_docx(self, file):
        doc = Document(file)
        return "\n".join([paragraph.text for paragraph in doc.paragraphs])

    def read_excel(self, file):
        wb = load_workbook(file)
        sheet = wb.active
        return "\n".join([" ".join([str(cell.value) for cell in row]) for row in sheet.rows])

    def extract_text_from_image(self, image):
        return pytesseract.image_to_string(image)